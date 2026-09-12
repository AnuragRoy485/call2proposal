from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from io import BytesIO
import re
from docx import Document
from fastapi.middleware.cors import CORSMiddleware
from .models import ProposalRequest, ProposalResponse, ExportRequest
from .engine import generate
from .providers import get_provider

app = FastAPI(title="Call2Proposal API", version="0.1.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["GET", "POST"], allow_headers=["*"])

@app.get("/health")
def health():
    try:
        provider = get_provider()
    except RuntimeError as exc:
        return {"status": "misconfigured", "error": str(exc)}
    return {"status": "ok", "provider": provider.name}

@app.post("/v1/proposals", response_model=ProposalResponse)
def create_proposal(request: ProposalRequest):
    return generate(request)


@app.post("/v1/export/docx")
def export_docx(request: ExportRequest):
    if not request.approved:
        raise HTTPException(status_code=409, detail="Human approval is required before export")
    doc = Document()
    for line in request.proposal_markdown.splitlines():
        text = line.strip()
        if not text:
            continue
        if text.startswith("# "):
            doc.add_heading(text[2:], level=0)
        elif text.startswith("## "):
            doc.add_heading(text[3:], level=1)
        elif re.match(r"^\d+\. ", text):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", text), style="List Number")
        elif text.startswith("- "):
            doc.add_paragraph(text[2:], style="List Bullet")
        else:
            doc.add_paragraph(text)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="proposal-{request.run_id}.docx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers=headers)
